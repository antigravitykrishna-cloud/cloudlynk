"""Builds Assignment3_Report.docx from outputs/ (run assignment3_mimic_eda.py first)."""
import json, pandas as pd
from docx import Document
from docx.shared import Inches, Pt

R = json.load(open("outputs/results.json")); O = "outputs/"
d = Document(); d.styles["Normal"].font.name = "Calibri"; d.styles["Normal"].font.size = Pt(11)

def H(t, l=1): d.add_heading(t, l)
def P(t, b=False):
    p = d.add_paragraph(); r = p.add_run(t); r.bold = b; return p
def IMG(f, w=6.4): d.add_picture(O + f, width=Inches(w))
def TAB(df):
    t = d.add_table(rows=1, cols=len(df.columns)); t.style = "Light Grid Accent 1"
    for i, c in enumerate(df.columns): t.rows[0].cells[i].text = str(c)
    for _, r in df.iterrows():
        cells = t.add_row().cells
        for i, x in enumerate(r): cells[i].text = str(x)
    d.add_paragraph()

c = R["cohort"]; a1 = R["A1"]; hr, cr = a1["HR"], a1["Creatinine"]
d.add_heading("Assignment 3: EDA Report + Standards Normalization on MIMIC-IV", 0)
P("Dataset: MIMIC-IV Clinical Database Demo v2.2 (100 patients). Code: assignment3_mimic_eda.py (all numbers/figures below are its output; log in outputs/run_log.txt).")
H("Cohort and method notes")
P(f"Cohort: ALL ICU stays in icustays: {c['stays']} stays, {c['admissions']} hospital admissions, {c['patients']} patients. "
  "Vitals come from chartevents (HR 220045, RR 220210, SpO2 220277, NIBP systolic 220179, NIBP mean 220181, Temp °C 223762 and Temp °F 223761 converted to °C). "
  "Labs come from labevents for the stay's hadm_id with charttime between ICU intime and outtime (creatinine 50912, BUN 51006, potassium 50971, sodium 50983, bicarbonate 50882, glucose 50931, hemoglobin 51222, WBC 51301, platelets 51265, lactate 50813).")
P("Two things to be upfront about:", True)
P("1. The physiologic range table from “Slide 55” was not among the files I was given, so I could not use it. I used my own conventional impossible-value limits (HR 20–300 bpm, SpO2 50–100 %, Temp 25–45 °C) and also report a stricter sensitivity range. The bounds are one dictionary (RANGES) in the script; if your slide differs, edit it and re-run and the counts in Part A3 update.")
P("2. The d_labitems.csv.gz in this demo has only the columns itemid, label, fluid, category. There is no loinc_code column, so the LOINC coverage answer in Part B is 0 % (see B3), not a number I could estimate from data.")

H("Part A: Exploratory Data Analysis")
H("A1. Distributions (10 pts): heart rate and creatinine", 2)
IMG("A1_distributions.png")
P(f"Heart rate (n={hr['n']:,} charted values): mean {hr['mean']:.1f}, median {hr['median']:.0f}, SD {hr['sd']:.1f}, skew {hr['skew']:.2f}, range {hr['min']:.0f}–{hr['max']:.0f}. "
  "It is unimodal and close to symmetric with a slightly longer right tail; the only odd feature is a handful of zeros (artefact, see A3). "
  f"Creatinine (n={cr['n']} values): median {cr['median']:.1f} mg/dL (IQR {cr['q1']:.1f}–{cr['q3']:.1f}), max {cr['max']:.1f}, skew {cr['skew']:.2f}, excess kurtosis {cr['kurtosis']:.1f}. "
  f"It is strongly right-skewed with a heavy tail (patients in renal failure); after a log transform the skew falls to {a1['log_Creatinine_skew']:.2f}, so it is roughly log-normal and a log transform is appropriate before using it in a linear or distance-based model. "
  "Note that the values are repeated measurements within stays, so they are not independent samples.")

H("A2. Missingness (10 pts)", 2)
P("Heatmap of 16 variables (6 vitals, 10 labs) across the 140 ICU stays. A cell is red if the variable was never measured during that stay (stay-level missingness, rows sorted by first care unit then length of stay), plus the percentage missing per variable.")
IMG("A2_missingness.png")
m = R["A2"]["missing_pct"]
P(f"Missing % by variable: lactate {m['Lactate']}%, platelets {m['Platelets']}%, WBC and hemoglobin {m['WBC']}%, chemistry panel (creatinine/BUN/glucose/bicarbonate) {m['Creatinine']}%, K/Na {m['Potassium']}%, MAP {m['MAP']}%, SBP and temperature {m['SBP']}%, HR/RR/SpO2 0%.")
IMG("A2_comissing_phi.png", 5)
cm = R["A2"]["top_comissing"]
P("Interpretation: this is not MCAR. (i) The missingness shows horizontal stripes: the same stays are missing whole groups of variables together (creatinine, BUN, glucose, bicarbonate have phi = 1.0 co-missingness; hemoglobin and WBC are missing in exactly the same stays, with platelets nearly so). That is a panel/ordering mechanism (a lab panel was either drawn or not), not independent random loss. "
  "(ii) Lactate is missing in 34 % of stays, far more than any other lab. Lactate is ordered when clinicians suspect shock or sepsis, so whether it is observed depends on the patient's (unobserved) severity, which is MNAR-like. "
  "(iii) The 3 stays with no chemistry panel in the ICU window are all very short (LOS 0.02–0.6 days, versus a median of about 2.2 days) and all ended in in-hospital death, so missingness is tied to outcome and stay length (MAR on LOS or MNAR); Fisher's exact test of chemistry-missing vs. death gives p = 0.0025, but with only 3 events this is suggestive, not conclusive. "
  "(iv) By contrast, missingness differs little by care unit (chi-square p > 0.1 for every variable, but most unit cells are tiny, so this test is weak). Vitals (HR, RR, SpO2) are essentially complete as they are charted continuously by monitors. "
  "Caveat: this is stay-level presence/absence; within-stay sampling frequency (e.g., labs drawn every 6 h vs. daily) is also informative and is a separate form of missingness.")

H("A3. Outliers (10 pts)", 2)
o = pd.DataFrame(R["A3"])
TAB(o[["variable", "unit", "valid_range", "n_obs", "n_flagged", "pct_flagged", "n_below", "n_above", "observed_min", "observed_max"]])
IMG("A3_outliers.png")
P("Range table used (assumption, see notes): HR 20–300, SpO2 50–100, Temp 25–45 °C (Fahrenheit readings converted to °C first). "
  "Flagged values: HR: three readings of exactly 0 (artefact/disconnected sensor or asystole not plausible alongside the rest of the record); SpO2: three readings below 50 % (29, 43 and 47 %); temperature: three readings of 97.2, 98.9 and 99.0 charted in the Celsius item, which are clearly Fahrenheit values entered in the wrong field (a unit-entry error). "
  f"Rates are tiny: {o.pct_flagged[0]:.3f}% (HR), {o.pct_flagged[1]:.3f}% (SpO2), {o.pct_flagged[2]:.3f}% (Temp). "
  f"With a stricter range (HR 30–220, SpO2 70–100, Temp 32–42 °C) the counts become {o.sens_n_flagged[0]}, {o.sens_n_flagged[1]} and {o.sens_n_flagged[2]} ({o.sens_pct[0]:.3f}%, {o.sens_pct[1]:.3f}%, {o.sens_pct[2]:.3f}%), so the result is sensitive to where the bounds are drawn, but remains well under 1 %. Flagged values are excluded in A5.")

H("A4. Temporal pattern (10 pts)", 2)
a4 = R["A4"]
IMG("A4_temporal.png", 6.6)
P(f"Patient chosen: subject_id {a4['subject_id']}, admission {a4['hadm_id']}, the admission with the most creatinine draws in the dataset ({a4['n_draws']} draws over {a4['span_days']} days). "
  f"In the first ~2 weeks draws are every ~6–12 h (median gap over the whole stay {a4['median_gap_h']} h, IQR {a4['gap_iqr'][0]}–{a4['gap_iqr'][1]} h); from about 10 February onward there is one draw a day at 04:15–07:35. "
  f"Across all creatinine draws in the demo, {a4['pct_draws_3_to_7am']}% fall between 03:00 and 06:59 (about 17 % would be expected if draws were uniform over 24 h), and this patient shows the same pattern ({a4['pct_this_pt_3_to_7am']}%). "
  "Interpretation: the sampling times look workflow-driven (a routine morning “AM labs” round) rather than driven by clinical events; the value trajectory itself (falling to 0.5 then rising to 4.1 mg/dL) is clinical, but when it is observed is set by hospital routine. The extra off-schedule draws early in the stay do reflect clinical attention. Practical consequence: timestamps/“time since last lab” features can leak workflow information, and time-aware models should account for the irregular, schedule-driven sampling. "
  "(Draws here include the whole hospital admission, not just ICU time.)")

H("A5. Correlation (10 pts)", 2)
IMG("A5_correlation.png", 6)
s = R["A5"]["strong"]
P(f"Heatmap of per-stay mean values for {R['A5']['n_features']} numeric features (pairwise-complete Pearson, outliers from A3 removed). Two pairs exceed |r| > 0.6: "
  + "; ".join(f"{x[0]}–{x[1]} r = {x[2]} (n = {x[3]} stays)" for x in s) + ". "
  "Creatinine–BUN (r = 0.72): both are renal-clearance markers that rise together when glomerular filtration falls (acute kidney injury/chronic kidney disease); BUN is additionally raised by dehydration, GI bleeding and catabolism. "
  "SBP–MAP (r = 0.77): mean arterial pressure is computed from systolic and diastolic pressure (MAP ≈ DBP + (SBP − DBP)/3), so the two are mechanically related. No other pair reaches 0.6, so redundancy is limited to these. These are stay-level correlations on a small sample (about 140), so they are descriptive, not causal.")

H("A6. Subgroup shift (10 pts): creatinine by gender", 2)
IMG("A6_subgroup.png", 5.6)
a6 = R["A6"]; vl, pl = a6["value_level"], a6["patient_level"]
P(f"All ICU creatinine measurements: women n = {vl['F']['n']}, median {vl['F']['median']} mg/dL (mean {vl['F']['mean']}); men n = {vl['M']['n']}, median {vl['M']['median']} (mean {vl['M']['mean']}); Mann–Whitney p = {a6['mwu_p_value_level']:.1e}. "
  f"Because measurements are clustered within patients, I repeated the test on per-patient means: women n = {pl['F']['n']}, median {pl['F']['median']}; men n = {pl['M']['n']}, median {pl['M']['median']}; p = {a6['mwu_p_patient_level']:.1e}. "
  "Men have clearly higher creatinine under both analyses. Part of this is physiological (higher muscle mass and creatinine production in men, so reference ranges differ by sex), and part may be case-mix (age, kidney disease, severity) that I did not adjust for. The practical point is that a model using raw creatinine sees a different distribution by sex, so a sex-specific reference or a fairness check is warranted. The small sample (100 patients) limits how far this generalizes.")

H("A7. Label distribution (10 pts)", 2)
IMG("A7_label.png", 3.4)
a7 = R["A7"]
P(f"hospital_expire_flag among the {c['admissions']} admissions with an ICU stay: {a7['cohort_admissions']['0']} survived (0) and {a7['cohort_admissions']['1']} died (1), a mortality of {a7['death_rate_pct']}% (imbalance ≈ {a7['imbalance']}:1). "
  f"At the stay level (a stay inherits its admission's flag) it is {a7['cohort_stays']['0']} vs {a7['cohort_stays']['1']}; I report admission-level numbers because stays from the same admission share a label. For all {sum(a7['all_admissions'].values())} admissions in the demo it is {a7['all_admissions']['0']} vs {a7['all_admissions']['1']} (the 15 deaths all occur in ICU admissions). "
  "The classes are imbalanced, so plain accuracy is misleading (always predicting survival gives 88 %); use AUROC/AUPRC, class weights or resampling, stratified splits, and note that with only 15 positives any model estimate will have very wide uncertainty. Patient-level (not row-level) splits are needed to avoid leakage.")

H("Part B: Standards Normalization")
b1 = R["B1"]
H("B1. Raw ICD codes vs. codes in ≥1 % of admissions (10 pts)", 2)
P(f"I used all {b1['n_admissions']} admissions that have diagnoses (4,506 diagnosis rows; requirement was 200+). Every code was found in d_icd_diagnoses (joined on icd_code + icd_version). ICD-9 and ICD-10 codes coexist in MIMIC-IV, so a code is identified as version:code.")
TAB(pd.DataFrame([["Unique raw ICD codes (version + code)", b1["unique_raw_codes"]],
                  ["(ignoring version, for reference)", b1["unique_raw_codes_ignoring_version"]],
                  [f"Unique codes in ≥1% of admissions (≥{b1['threshold_admissions']} of {b1['n_admissions']}, since 1% of 275 = 2.75)", b1["codes_ge_1pct"]],
                  ["Share of unique codes that reach 1%", f"{100*b1['codes_ge_1pct']/b1['unique_raw_codes']:.1f}%"]], columns=["Metric", "Value"]))
IMG("B1_long_tail.png", 4.8)
P("Most-prevalent codes (note that the same concept appears under both ICD versions, e.g. hypertension 4019 vs I10, hyperlipidemia 2724 vs E785):")
TAB(pd.DataFrame([[x[0], x[1], x[2], x[3][:70]] for x in b1["top"]], columns=["code", "admissions", "%", "description"]))
P(f"Interpretation: the vocabulary is a long tail: {b1['unique_raw_codes']} distinct codes but only {b1['codes_ge_1pct']} appear in at least 1% of admissions, and the most common code covers only {b1['max_prevalence_pct']}% of admissions. Most codes occur in one or two admissions, which is too few to learn from, motivating rollup and frequency thresholds. (In this 100-patient demo the 1% cut-off is only 3 admissions, so the count would be much smaller in full MIMIC-IV.)")

b2 = R["B2"]
H("B2. Roll-up to 3-character ICD families (5 pts)", 2)
TAB(pd.read_csv(O + "B2_rollup_top20.csv").rename(columns={"raw_code": "raw code (version:code)", "family": "3-char family"}))
P(f"The 20 most frequent codes above map to {b2['top20_unique_families']} distinct families (none collapse because these frequent codes already belong to different families). Applied to ALL {b2['unique_raw']} unique raw codes, the roll-up leaves {b2['unique_families']} unique categories, a {b2['reduction_pct']}% reduction; {b2['families_ge_1pct']} of these families occur in ≥1% of admissions (versus {b1['codes_ge_1pct']} raw codes), and {b2['families_with_multiple_raw']} families merge two or more raw codes. "
  "Caveats: (1) the roll-up is done within each ICD version, so equivalent ICD-9 and ICD-10 concepts (9:401 vs 10:I10) stay separate until a cross-version map (e.g., GEMs) is applied; (2) ICD-9 E-codes properly have a 4-character family, and a plain 3-character cut is coarser for them; (3) MIMIC stores codes without the dot, so J18.9 appears as J189 and rolls up to J18.")

b3 = R["B3"]
H("B3. LOINC coverage of lab itemids (5 pts)", 2)
P(f"The dataset has {b3['n_itemids_in_labevents']} distinct itemids in labevents, all present in d_labitems. The d_labitems table in this download has the columns {', '.join(b3['d_labitems_columns'])}, with no loinc_code column. "
  f"Therefore the fraction of itemids with a non-null LOINC code is {b3['n_with_loinc']}/{b3['n_itemids_in_labevents']} = {b3['fraction_with_loinc']*100:.0f}%. I did not find LOINC codes anywhere else in the download (a text search of the package for “loinc” returned nothing). "
  "If your course materials show a loinc_code column, it comes from a different release of d_labitems; the script handles this: if the column exists it computes the true non-null fraction.")
P(f"For scale: the itemids split into fluid Blood {b3['by_fluid']['Blood']}, Urine {b3['by_fluid']['Urine']} and others, and categories Hematology {b3['by_category']['Hematology']}, Chemistry {b3['by_category']['Chemistry']}, Blood Gas {b3['by_category']['Blood Gas']}. The top 25 itemids account for {b3['event_share_top']['25']*100:.0f}% of lab events and the top 50 for {b3['event_share_top']['50']*100:.0f}%, so manual curation of the head is cheap and covers most of the data. {b3['duplicate_labels_across_itemids']} lower-cased labels are shared by more than one itemid (usually the same test with a different fluid or analyzer).")

H("B4. Fallback strategy for itemids without a LOINC code (10 pts)", 2)
for t in [
 "1. Curate the head by hand. Rank itemids by event count and have a clinician or informatician map the top ~50–100 to LOINC (this covers about 79–94% of events here). Store the result as a reviewed crosswalk table (itemid → LOINC, mapping method, confidence, reviewer).",
 "2. Normalize label strings for the tail. Lower-case, strip punctuation and extra whitespace, expand abbreviations (e.g., “Urea Nitrogen” = BUN, “WBC” = white blood cell count, “Hct” = hematocrit), and remove assay modifiers. Then exact-match against LOINC long common names and the LOINC “component” field.",
 "3. Use more than the name. A LOINC code is defined by component, property, time, system (specimen), scale and method, so build the query from several fields: label → component; the fluid column (Blood, Urine, CSF, Pleural...) → system; valueuom (mg/dL, mEq/L, %, #/uL) → property type and unit; category (Chemistry/Hematology/Blood Gas) as a sanity check; and the reference range (ref_range_lower/upper) as a plausibility check. Two itemids with the same label but different fluid must not map to the same code.",
 "4. Fuzzy and embedding matching with human review. For labels without an exact match, rank LOINC candidates by token-set or character n-gram similarity (e.g., TF-IDF cosine, Jaro-Winkler) and optionally by embeddings from a biomedical language model (SapBERT or similar); accept automatically only above a high threshold and when fluid and unit agree, and queue the rest for review.",
 "5. Use value-distribution evidence. Compare the itemid's observed unit and value distribution with that of a confidently mapped item (for example, a sodium-like range around 135–145 mmol/L supports a sodium mapping) to detect and fix wrong matches.",
 "6. Fail safe. Items that remain unmatched keep their local itemid as a namespaced code (e.g., MIMIC:51234) rather than being dropped, are flagged “unmapped”, and are excluded from or handled separately in cross-site models. Validate the pipeline by measuring precision on a held-out, hand-labelled sample of mapped items, and version the crosswalk so the mapping is reproducible."]:
    P(t)

H("Reproducibility")
P("Run: python assignment3_mimic_eda.py <path to mimic-iv-clinical-database-demo-2.2> outputs, then python build_report.py. Requires pandas, numpy, scipy, matplotlib, seaborn, python-docx. All figures in this report are generated by that script.")
d.save("Assignment3_Report.docx")
